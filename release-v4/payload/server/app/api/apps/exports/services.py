import os
from io import BytesIO

from django.conf import settings
from django.db.models import Q
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

from apps.common.utils import normalize_search_text
from apps.profiles.models import Profile


EXPORT_HEADERS = [
    "STT",
    "Mã hồ sơ",
    "Tên hồ sơ",
    "Năm",
    "Phông",
    "Mục lục",
    "Kho",
    "Vị trí/Kệ",
    "Hộp số",
    "Tệp số",
    "Số trang",
    "Thời hạn bảo quản",
    "Ghi chú",
]


def get_export_profiles_queryset(params):
    """
    Lấy danh sách hồ sơ theo điều kiện export.

    Hỗ trợ query params:
    - q
    - year
    - box_number
    - file_number
    - fond_id
    - catalog_id
    """

    queryset = (
        Profile.objects
        .select_related(
            "catalog",
            "catalog__fond",
            "storage_file",
            "storage_file__box",
            "storage_file__box__location",
            "storage_file__box__location__warehouse",
        )
        .filter(is_deleted=False)
        .order_by("catalog__fond__code", "year", "profile_code")
    )

    q = params.get("q")
    year = params.get("year")
    box_number = params.get("box_number")
    file_number = params.get("file_number")
    fond_id = params.get("fond_id")
    catalog_id = params.get("catalog_id")

    if q:
        normalized_q = normalize_search_text(q)

        queryset = queryset.filter(
            Q(search_text__icontains=normalized_q)
            | Q(profile_code__icontains=q)
            | Q(title__icontains=q)
            | Q(description__icontains=q)
        )

    if year:
        queryset = queryset.filter(year=year)

    if box_number:
        queryset = queryset.filter(
            storage_file__box__box_number__icontains=box_number
        )

    if file_number:
        queryset = queryset.filter(
            storage_file__file_number__icontains=file_number
        )

    if fond_id:
        queryset = queryset.filter(catalog__fond_id=fond_id)

    if catalog_id:
        queryset = queryset.filter(catalog_id=catalog_id)

    return queryset


def profile_to_export_row(index, profile):
    """
    Chuyển 1 hồ sơ thành 1 dòng export.
    """

    catalog = profile.catalog
    fond = catalog.fond if catalog else None
    storage_file = profile.storage_file
    box = storage_file.box if storage_file else None
    location = box.location if box else None
    warehouse = location.warehouse if location else None

    warehouse_text = ""
    if warehouse:
        warehouse_text = f"{warehouse.code} - {warehouse.name}"

    location_text = ""
    if location:
        location_text = f"{location.code} - {location.name}"

    return [
        index,
        profile.profile_code or "",
        profile.title or "",
        profile.year or "",
        f"{fond.code} - {fond.name}" if fond else "",
        f"{catalog.code} - {catalog.name}" if catalog else "",
        warehouse_text,
        location_text,
        box.box_number if box else "",
        storage_file.file_number if storage_file else "",
        profile.total_pages or "",
        profile.retention_period or "",
        profile.note or "",
    ]


def build_profiles_excel_export(params):
    """
    Tạo file Excel export danh sách hồ sơ.
    """

    queryset = get_export_profiles_queryset(params)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Danh sách hồ sơ"

    title = "DANH SÁCH HỒ SƠ LƯU TRỮ"
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(EXPORT_HEADERS))
    title_cell = worksheet.cell(row=1, column=1, value=title)
    title_cell.font = Font(bold=True, size=14)
    title_cell.alignment = Alignment(horizontal="center")

    worksheet.append([])
    worksheet.append(EXPORT_HEADERS)

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(bold=True)

    for cell in worksheet[3]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for index, profile in enumerate(queryset, start=1):
        worksheet.append(profile_to_export_row(index, profile))

    for row in worksheet.iter_rows(min_row=4):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    widths = [8, 24, 45, 10, 35, 35, 28, 28, 14, 14, 12, 20, 35]

    for col_index, width in enumerate(widths, start=1):
        column_letter = get_column_letter(col_index)
        worksheet.column_dimensions[column_letter].width = width

    worksheet.freeze_panes = "A4"
    worksheet.auto_filter.ref = f"A3:{get_column_letter(len(EXPORT_HEADERS))}{worksheet.max_row}"

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return output, queryset.count()


def register_pdf_font():
    """
    Đăng ký font hỗ trợ tiếng Việt cho PDF.

    Trên Windows ưu tiên Arial.
    Nếu không tìm thấy thì dùng Helvetica mặc định.
    """

    font_candidates = [
        r"C:\Windows\Fonts\arial.ttf",
        r"C:\Windows\Fonts\Arial.ttf",
        os.path.join(settings.BASE_DIR, "fonts", "arial.ttf"),
        os.path.join(settings.BASE_DIR, "fonts", "DejaVuSans.ttf"),
    ]

    for font_path in font_candidates:
        if os.path.exists(font_path):
            pdfmetrics.registerFont(TTFont("VietnameseFont", font_path))
            return "VietnameseFont"

    return "Helvetica"


def as_paragraph(value, style):
    """
    Chuyển text sang Paragraph để PDF tự xuống dòng.
    """

    value = "" if value is None else str(value)
    value = (
        value.replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )
    return Paragraph(value, style)


def build_profiles_pdf_export(params):
    """
    Tạo file PDF export danh sách hồ sơ.
    """

    queryset = get_export_profiles_queryset(params)
    font_name = register_pdf_font()

    output = BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=10 * mm,
        leftMargin=10 * mm,
        topMargin=10 * mm,
        bottomMargin=10 * mm,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "VietnameseTitle",
        parent=styles["Title"],
        fontName=font_name,
        fontSize=15,
        leading=18,
        alignment=TA_CENTER,
        spaceAfter=8,
    )

    normal_style = ParagraphStyle(
        "VietnameseNormal",
        parent=styles["Normal"],
        fontName=font_name,
        fontSize=7,
        leading=9,
    )

    header_style = ParagraphStyle(
        "VietnameseHeader",
        parent=styles["Normal"],
        fontName=font_name,
        fontSize=7,
        leading=9,
        alignment=TA_CENTER,
    )

    elements = []

    elements.append(Paragraph("DANH SÁCH HỒ SƠ LƯU TRỮ", title_style))
    elements.append(Spacer(1, 5))

    table_data = [
        [as_paragraph(header, header_style) for header in EXPORT_HEADERS]
    ]

    for index, profile in enumerate(queryset, start=1):
        row = profile_to_export_row(index, profile)
        table_data.append([as_paragraph(value, normal_style) for value in row])

    if queryset.count() == 0:
        table_data.append(
            [
                as_paragraph("Không có dữ liệu", normal_style),
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
                "",
            ]
        )

    column_widths = [
        10 * mm,
        26 * mm,
        45 * mm,
        12 * mm,
        34 * mm,
        34 * mm,
        28 * mm,
        26 * mm,
        14 * mm,
        14 * mm,
        14 * mm,
        22 * mm,
        34 * mm,
    ]

    table = Table(
        table_data,
        colWidths=column_widths,
        repeatRows=1,
    )

    table.setStyle(
        TableStyle(
            [
                ("FONTNAME", (0, 0), (-1, -1), font_name),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9EAF7")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.black),
                ("ALIGN", (0, 0), (-1, 0), "CENTER"),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )

    elements.append(table)

    document.build(elements)

    output.seek(0)

    return output, queryset.count()