from datetime import date, datetime
from io import BytesIO
import re

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from apps.catalogs.models import Catalog
from apps.documents.models import Document
from apps.fonds.models import Fond
from apps.profiles.models import Profile
from apps.storage.models import StorageBox, StorageFile, StorageLocation, Warehouse


IMPORT_COLUMNS = [
    "fond_code",
    "fond_name",
    "catalog_code",
    "catalog_name",
    "year",
    "warehouse_code",
    "warehouse_name",
    "location_code",
    "location_name",
    "box_number",
    "box_title",
    "file_number",
    "file_title",
    "profile_code",
    "profile_title",
    "profile_description",
    "total_pages",
    "retention_period",
    "language",
    "note",
    "document_code",
    "document_title",
    "document_date",
    "author",
    "page_start",
    "page_end",
    "summary",
]

REQUIRED_COLUMNS = [
    "fond_code",
    "fond_name",
    "catalog_code",
    "catalog_name",
    "year",
    "warehouse_code",
    "warehouse_name",
    "location_code",
    "location_name",
    "box_number",
    "file_number",
    "profile_code",
    "profile_title",
]

SAMPLE_ROW = {
    "fond_code": "UBND-XA-A",
    "fond_name": "Ủy ban nhân dân xã A",
    "catalog_code": "ML-2024",
    "catalog_name": "Mục lục hồ sơ năm 2024",
    "year": 2024,
    "warehouse_code": "KHO-01",
    "warehouse_name": "Kho lưu trữ số 01",
    "location_code": "KE-A",
    "location_name": "Kệ A",
    "box_number": "12",
    "box_title": "Hộp hồ sơ số 12",
    "file_number": "03",
    "file_title": "Tệp hồ sơ số 03",
    "profile_code": "235/BC-UBND",
    "profile_title": "Báo cáo công tác dân tộc năm 2024",
    "profile_description": "Hồ sơ báo cáo công tác dân tộc của UBND xã A",
    "total_pages": 25,
    "retention_period": "Vĩnh viễn",
    "language": "Tiếng Việt",
    "note": "Dòng mẫu import",
    "document_code": "235/BC-UBND",
    "document_title": "Báo cáo công tác dân tộc",
    "document_date": "2024-12-20",
    "author": "UBND xã A",
    "page_start": 1,
    "page_end": 25,
    "summary": "Báo cáo kết quả công tác dân tộc năm 2024",
}


def clean_value(value):
    """
    Chuẩn hóa giá trị đọc từ Excel.

    - None giữ nguyên None.
    - Chuỗi thì strip khoảng trắng đầu/cuối.
    - Các kiểu khác giữ nguyên.
    """

    if value is None:
        return None

    if isinstance(value, str):
        value = value.strip()
        return value if value != "" else None

    return value


def to_str(value):
    """
    Ép giá trị về chuỗi để lưu các mã/ký hiệu.

    Ví dụ:
    Excel nhập số 12 thì chuyển thành "12".
    """

    value = clean_value(value)

    if value is None:
        return ""

    if isinstance(value, float) and value.is_integer():
        return str(int(value))

    return str(value).strip()


def to_int(value):
    """
    Ép giá trị Excel về số nguyên.

    Dùng cho year, total_pages, page_start, page_end.
    """

    value = clean_value(value)

    if value is None:
        return None

    if isinstance(value, int):
        return value

    if isinstance(value, float):
        return int(value)

    try:
        return int(str(value).strip())
    except ValueError:
        return None


def to_date(value):
    """
    Ép giá trị Excel về date.

    Hỗ trợ:
    - Date cell của Excel
    - yyyy-mm-dd
    - dd/mm/yyyy
    """

    value = clean_value(value)

    if value is None:
        return None

    if isinstance(value, date) and not isinstance(value, datetime):
        return value

    if isinstance(value, datetime):
        return value.date()

    value = str(value).strip()

    for fmt in ["%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"]:
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            continue

    return None


def build_profiles_import_template():
    """
    Tạo file Excel mẫu import hồ sơ.
    """

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "profiles_import"

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    required_fill = PatternFill("solid", fgColor="FFD966")
    header_font = Font(bold=True)

    for col_index, column_name in enumerate(IMPORT_COLUMNS, start=1):
        cell = worksheet.cell(row=1, column=col_index, value=column_name)
        cell.font = header_font

        if column_name in REQUIRED_COLUMNS:
            cell.fill = required_fill
        else:
            cell.fill = header_fill

    for col_index, column_name in enumerate(IMPORT_COLUMNS, start=1):
        worksheet.cell(row=2, column=col_index, value=SAMPLE_ROW.get(column_name))

    for col_index, column_name in enumerate(IMPORT_COLUMNS, start=1):
        column_letter = get_column_letter(col_index)
        width = max(len(column_name) + 4, 18)
        worksheet.column_dimensions[column_letter].width = width

    worksheet.freeze_panes = "A2"

    notes_sheet = workbook.create_sheet("huong_dan")
    notes = [
        ["Cột màu vàng", "Bắt buộc nhập"],
        ["Cột màu xanh", "Không bắt buộc"],
        ["document_code/document_title", "Nếu có thì hệ thống sẽ tạo/cập nhật thành phần hồ sơ"],
        ["document_date", "Định dạng khuyến nghị: yyyy-mm-dd, ví dụ 2024-12-20"],
        ["dry_run=true", "Chỉ kiểm tra dữ liệu, không ghi database"],
        ["dry_run=false", "Import thật vào database"],
    ]

    for row in notes:
        notes_sheet.append(row)

    notes_sheet.column_dimensions["A"].width = 28
    notes_sheet.column_dimensions["B"].width = 80

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return output


def read_excel_rows(uploaded_file, sheet_name=None):
    """
    Đọc file Excel upload và chuyển thành list dict.

    Dòng 1 là header.
    Từ dòng 2 trở đi là dữ liệu.
    """

    workbook = load_workbook(uploaded_file, data_only=True)

    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            return [], [
                {
                    "row": 1,
                    "field": "sheet_name",
                    "message": f'Không tìm thấy sheet "{sheet_name}" trong file Excel.',
                }
            ]

        worksheet = workbook[sheet_name]
    else:
        worksheet = workbook.active

    header_row = [
        clean_value(cell.value)
        for cell in worksheet[1]
    ]

    headers = [str(header).strip() if header else "" for header in header_row]

    missing_columns = [
        column for column in REQUIRED_COLUMNS
        if column not in headers
    ]

    if missing_columns:
        return [], [
            {
                "row": 1,
                "field": "header",
                "message": f"Thiếu cột bắt buộc: {', '.join(missing_columns)}",
            }
        ]

    rows = []

    for row_index, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        raw = dict(zip(headers, row))

        if all(clean_value(value) is None for value in raw.values()):
            continue

        item = {
            column: clean_value(raw.get(column))
            for column in IMPORT_COLUMNS
        }
        item["_row_number"] = row_index
        item["_sheet_name"] = worksheet.title
        rows.append(item)

    return rows, []


def validate_import_row(row):
    """
    Validate một dòng dữ liệu.
    """

    errors = []
    row_number = row.get("_row_number")

    for column in REQUIRED_COLUMNS:
        if clean_value(row.get(column)) is None:
            errors.append(
                {
                    "row": row_number,
                    "field": column,
                    "message": "Không được để trống.",
                }
            )

    year = to_int(row.get("year"))

    if year is None:
        errors.append(
            {
                "row": row_number,
                "field": "year",
                "message": "Năm không hợp lệ.",
            }
        )

    total_pages = to_int(row.get("total_pages"))

    if row.get("total_pages") is not None and total_pages is None:
        errors.append(
            {
                "row": row_number,
                "field": "total_pages",
                "message": "Số trang không hợp lệ.",
            }
        )

    page_start = to_int(row.get("page_start"))
    page_end = to_int(row.get("page_end"))

    if row.get("page_start") is not None and page_start is None:
        errors.append(
            {
                "row": row_number,
                "field": "page_start",
                "message": "Trang bắt đầu không hợp lệ.",
            }
        )

    if row.get("page_end") is not None and page_end is None:
        errors.append(
            {
                "row": row_number,
                "field": "page_end",
                "message": "Trang kết thúc không hợp lệ.",
            }
        )

    if page_start is not None and page_end is not None and page_end < page_start:
        errors.append(
            {
                "row": row_number,
                "field": "page_end",
                "message": "Trang kết thúc không được nhỏ hơn trang bắt đầu.",
            }
        )

    if row.get("document_date") is not None and to_date(row.get("document_date")) is None:
        errors.append(
            {
                "row": row_number,
                "field": "document_date",
                "message": "Ngày văn bản không hợp lệ. Dùng yyyy-mm-dd hoặc dd/mm/yyyy.",
            }
        )

    return errors


def validate_import_rows(rows):
    """
    Validate toàn bộ file Excel.
    """

    errors = []

    for row in rows:
        errors.extend(validate_import_row(row))

    return errors


def import_one_row(row):
    """
    Import một dòng Excel vào database.

    Dùng update_or_create để:
    - Nếu dữ liệu chưa có thì tạo mới.
    - Nếu dữ liệu đã có thì cập nhật.
    """

    year = to_int(row.get("year"))

    fond, fond_created = Fond.objects.update_or_create(
        code=to_str(row.get("fond_code")),
        defaults={
            "name": to_str(row.get("fond_name")),
            "description": "",
        },
    )

    catalog, catalog_created = Catalog.objects.update_or_create(
        fond=fond,
        code=to_str(row.get("catalog_code")),
        defaults={
            "name": to_str(row.get("catalog_name")),
            "year": year,
            "description": "",
        },
    )

    warehouse, warehouse_created = Warehouse.objects.update_or_create(
        code=to_str(row.get("warehouse_code")),
        defaults={
            "name": to_str(row.get("warehouse_name")),
            "address": "",
            "description": "",
        },
    )

    location, location_created = StorageLocation.objects.update_or_create(
        warehouse=warehouse,
        code=to_str(row.get("location_code")),
        defaults={
            "name": to_str(row.get("location_name")),
            "description": "",
        },
    )

    box_title = to_str(row.get("box_title")) or f"Hộp số {to_str(row.get('box_number'))}"

    box, box_created = StorageBox.objects.update_or_create(
        location=location,
        box_number=to_str(row.get("box_number")),
        defaults={
            "fond": fond,
            "catalog": catalog,
            "title": box_title,
            "description": "",
        },
    )

    file_title = to_str(row.get("file_title")) or f"Tệp số {to_str(row.get('file_number'))}"

    storage_file, storage_file_created = StorageFile.objects.update_or_create(
        box=box,
        file_number=to_str(row.get("file_number")),
        defaults={
            "title": file_title,
            "description": "",
        },
    )

    profile, profile_created = Profile.objects.update_or_create(
        profile_code=to_str(row.get("profile_code")),
        defaults={
            "catalog": catalog,
            "storage_file": storage_file,
            "title": to_str(row.get("profile_title")),
            "description": to_str(row.get("profile_description")),
            "year": year,
            "total_pages": to_int(row.get("total_pages")),
            "retention_period": to_str(row.get("retention_period")),
            "language": to_str(row.get("language")) or "Tiếng Việt",
            "note": to_str(row.get("note")),
        },
    )

    document = None
    document_created = False

    document_code = to_str(row.get("document_code"))
    document_title = to_str(row.get("document_title"))

    if document_code or document_title:
        document_lookup_code = document_code or f"{profile.profile_code}-DOC"

        document, document_created = Document.objects.update_or_create(
            profile=profile,
            document_code=document_lookup_code,
            defaults={
                "title": document_title or document_lookup_code,
                "document_date": to_date(row.get("document_date")),
                "author": to_str(row.get("author")),
                "page_start": to_int(row.get("page_start")),
                "page_end": to_int(row.get("page_end")),
                "summary": to_str(row.get("summary")),
            },
        )

    return {
        "fond_created": fond_created,
        "catalog_created": catalog_created,
        "warehouse_created": warehouse_created,
        "location_created": location_created,
        "box_created": box_created,
        "storage_file_created": storage_file_created,
        "profile_created": profile_created,
        "document_created": document_created,
        "profile_id": profile.id,
        "profile_code": profile.profile_code,
        "document_id": document.id if document else None,
    }


@transaction.atomic
def import_profiles_from_excel(uploaded_file, dry_run=False, sheet_name=None):
    """
    Import hồ sơ từ Excel.

    Nếu dry_run=True:
    - chỉ đọc và validate
    - không ghi database

    Nếu dry_run=False:
    - validate trước
    - nếu không lỗi thì import thật
    """

    rows, header_errors = read_excel_rows(uploaded_file, sheet_name=sheet_name)

    if header_errors:
        return {
            "dry_run": dry_run,
            "total_rows": 0,
            "success_count": 0,
            "error_count": len(header_errors),
            "errors": header_errors,
            "items": [],
        }

    validation_errors = validate_import_rows(rows)

    if validation_errors:
        return {
            "dry_run": dry_run,
            "total_rows": len(rows),
            "success_count": 0,
            "error_count": len(validation_errors),
            "errors": validation_errors,
            "items": [],
        }

    if dry_run:
        return {
            "dry_run": True,
            "total_rows": len(rows),
            "success_count": len(rows),
            "error_count": 0,
            "errors": [],
            "items": [
                {
                    "row": row.get("_row_number"),
                    "sheet": row.get("_sheet_name"),
                    "profile_code": to_str(row.get("profile_code")),
                    "profile_title": to_str(row.get("profile_title")),
                    "message": "Dữ liệu hợp lệ, chưa ghi database.",
                }
                for row in rows
            ],
        }

    imported_items = []

    for row in rows:
        result = import_one_row(row)
        result["row"] = row.get("_row_number")
        result["sheet"] = row.get("_sheet_name")
        imported_items.append(result)

    return {
        "dry_run": False,
        "total_rows": len(rows),
        "success_count": len(imported_items),
        "error_count": 0,
        "errors": [],
        "items": imported_items,
    }


# STEP 21.13C2G: Import hồ sơ từ danh sách đã trích xuất trên web.
def _c2g_to_text(value):
    if value is None:
        return ""
    return str(value).strip()


def _c2g_to_int(value):
    text = _c2g_to_text(value)
    if not text:
        return None

    match = re.search(r"\d+", text.replace(".", "").replace(",", ""))
    if not match:
        return None

    try:
        return int(match.group(0))
    except ValueError:
        return None


def _c2g_extract_year(value):
    text = _c2g_to_text(value)
    matches = re.findall(r"(19\d{2}|20\d{2})", text)

    if not matches:
        return None

    try:
        return int(matches[0])
    except ValueError:
        return None


def _c2g_get_default_catalog():
    return (
        Catalog.objects
        .filter(is_deleted=False)
        .order_by("id")
        .first()
    )


def _c2g_get_default_storage_file():
    return (
        StorageFile.objects
        .filter(is_deleted=False)
        .order_by("id")
        .first()
    )




# STEP 21.13C2H: Chọn vị trí lưu khi import từ parsedRows.
def _c2h_get_catalog_by_id(value):
    text = _c2g_to_text(value)

    if not text:
        return None

    try:
        return Catalog.objects.filter(id=int(text), is_deleted=False).first()
    except (TypeError, ValueError):
        return None


def _c2h_get_storage_file_by_id(value):
    text = _c2g_to_text(value)

    if not text:
        return None

    try:
        return StorageFile.objects.filter(id=int(text), is_deleted=False).first()
    except (TypeError, ValueError):
        return None


def import_profiles_from_client_rows(client_rows, dry_run=False):
    """
    Import trực tiếp từ parsedRows frontend đã trích xuất trong popup Excel.

    Dùng cho file mục lục hồ sơ thực tế không theo template backend cũ.
    Backend chỉ tạo/cập nhật hồ sơ; không tạo bản ghi ở bước này.
    """
    rows = client_rows or []
    errors = []

    if not isinstance(rows, list):
        return {
            "dry_run": dry_run,
            "source": "client_rows",
            "total_rows": 0,
            "success_count": 0,
            "error_count": 1,
            "errors": [
                {
                    "row": 1,
                    "field": "client_rows_json",
                    "message": "Danh sách hồ sơ trích xuất không hợp lệ.",
                }
            ],
            "items": [],
        }

    catalog = _c2g_get_default_catalog()
    storage_file = _c2g_get_default_storage_file()

    # Có thể dùng default nếu người dùng chưa chọn, nhưng ưu tiên vị trí từng dòng.
    # Nếu không có default và dòng không gửi catalog_id/storage_file_id thì sẽ báo lỗi theo dòng.

    normalized_rows = []

    for index, raw in enumerate(rows, start=1):
        if not isinstance(raw, dict):
            errors.append({
                "row": index,
                "field": "row",
                "message": "Dòng trích xuất không hợp lệ.",
            })
            continue

        row_number = raw.get("rowNumber") or raw.get("row_number") or index
        profile_code = _c2g_to_text(raw.get("profileCode") or raw.get("profile_code"))
        title = _c2g_to_text(raw.get("profileTitle") or raw.get("profile_title"))
        date_range = _c2g_to_text(raw.get("dateRange") or raw.get("date_range"))
        total_pages = _c2g_to_int(raw.get("totalPages") or raw.get("total_pages"))
        retention_period = _c2g_to_text(raw.get("retentionPeriod") or raw.get("retention_period"))
        note = _c2g_to_text(raw.get("notes") or raw.get("note"))
        sheet_name = _c2g_to_text(raw.get("sheetName") or raw.get("sheet_name"))
        catalog_id = _c2g_to_text(raw.get("catalog_id") or raw.get("catalogId"))
        storage_file_id = _c2g_to_text(raw.get("storage_file_id") or raw.get("storageFileId"))

        row_catalog = _c2h_get_catalog_by_id(catalog_id) if catalog_id else catalog
        row_storage_file = _c2h_get_storage_file_by_id(storage_file_id) if storage_file_id else storage_file

        if not profile_code:
            errors.append({
                "sheet": sheet_name,
                "row": row_number,
                "field": "profile_code",
                "message": "Thiếu số, ký hiệu hồ sơ.",
            })
            continue

        if row_catalog is None:
            errors.append({
                "sheet": sheet_name,
                "row": row_number,
                "field": "catalog_id",
                "message": "Chưa chọn Mục lục cho hồ sơ này.",
            })
            continue

        if row_storage_file is None:
            errors.append({
                "sheet": sheet_name,
                "row": row_number,
                "field": "storage_file_id",
                "message": "Chưa chọn Tệp lưu trữ cho hồ sơ này.",
            })
            continue

        if not title:
            errors.append({
                "sheet": sheet_name,
                "row": row_number,
                "field": "title",
                "message": "Thiếu tên hồ sơ.",
            })
            continue

        normalized_rows.append({
            "sheet": sheet_name,
            "row": row_number,
            "profile_code": profile_code,
            "catalog": row_catalog,
            "storage_file": row_storage_file,
            "title": title,
            "description": date_range,
            "year": _c2g_extract_year(date_range),
            "total_pages": total_pages,
            "retention_period": retention_period or "20 năm",
            "language": "vi",
            "note": note,
        })

    if errors:
        return {
            "dry_run": dry_run,
            "source": "client_rows",
            "total_rows": len(rows),
            "success_count": 0,
            "error_count": len(errors),
            "errors": errors,
            "items": [],
        }

    if dry_run:
        return {
            "dry_run": True,
            "source": "client_rows",
            "total_rows": len(normalized_rows),
            "success_count": len(normalized_rows),
            "error_count": 0,
            "errors": [],
            "items": [
                {
                    "sheet": row["sheet"],
                    "row": row["row"],
                    "profile_code": row["profile_code"],
                    "profile_title": row["title"],
                    "catalog_id": row["catalog"].id,
                    "catalog_name": getattr(row["catalog"], "name", ""),
                    "storage_file_id": row["storage_file"].id,
                    "storage_file_number": getattr(row["storage_file"], "file_number", ""),
                    "message": "Dữ liệu đã trích xuất hợp lệ, chưa ghi database.",
                }
                for row in normalized_rows
            ],
        }

    imported_items = []

    with transaction.atomic():
        for row in normalized_rows:
            profile = (
                Profile.objects
                .filter(profile_code=row["profile_code"], is_deleted=False)
                .first()
            )

            created = False

            if profile is None:
                profile = Profile(profile_code=row["profile_code"])
                created = True

            profile.catalog = row["catalog"]
            profile.storage_file = row["storage_file"]
            profile.title = row["title"]
            profile.description = row["description"]
            profile.year = row["year"]
            profile.total_pages = row["total_pages"] or 0
            profile.retention_period = row["retention_period"]
            profile.language = row["language"]
            profile.note = row["note"]
            profile.is_deleted = False
            profile.save()

            imported_items.append({
                "sheet": row["sheet"],
                "row": row["row"],
                "profile_id": profile.id,
                "profile_code": profile.profile_code,
                "profile_title": profile.title,
                "status": "created" if created else "updated",
                "message": "Đã tạo hồ sơ." if created else "Đã cập nhật hồ sơ.",
            })

    return {
        "dry_run": False,
        "source": "client_rows",
        "total_rows": len(normalized_rows),
        "success_count": len(imported_items),
        "error_count": 0,
        "errors": [],
        "items": imported_items,
    }

